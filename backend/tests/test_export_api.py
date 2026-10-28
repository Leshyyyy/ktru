"""GET /jobs/{id}/export (FR-7.1, FR-7.2): форматы, коды в строках исходного файла, X-Unreviewed-Count, статусы."""

from __future__ import annotations

import csv
import io
import json
import uuid
from typing import Any

import openpyxl
import pytest

from .conftest import HEADER, ROWS, Stack, xlsx_bytes
from .openapi import assert_matches_schema

pytestmark = pytest.mark.db

ADDED = ["Код КТРУ", "Наименование КТРУ", "Уверенность (%)", "Статус", "Источник", "Предупреждения"]


def _export(stack: Stack, job_id: str, **params: Any) -> Any:
    return stack.client.get(f"/api/v1/jobs/{job_id}/export", params=params)


def _error(response: Any, status: int, code: str) -> dict[str, Any]:
    assert response.status_code == status, response.text
    body = response.json()
    assert_matches_schema("Error", body)
    assert body["error"]["code"] == code
    return body["error"]


def _job(stack: Stack, table: list[list[Any]], filename: str = "spec.xlsx") -> dict[str, Any]:
    """Как фронт: строка заголовка — та, что нашёл POST /files."""
    info = stack.upload(xlsx_bytes(table), filename=filename)
    job = stack.create_job(info["file_id"], header_row=info["sheets"][0]["header_row"])
    stack.drain()
    return stack.client.get(f"/api/v1/jobs/{job['job_id']}").json()


def test_codes_land_in_their_rows_of_the_source_file(stack: Stack) -> None:
    table = [
        ["Спецификация на поставку офисной техники", None, None, None],
        HEADER,
        ROWS[0],
        ROWS[1],
        [None, None, None, None],  # пустая строка в середине таблицы
        ROWS[2],
        ["Итого", None, 55, None],
    ]
    job = _job(stack, table, filename="Спецификация №1.xlsx")
    assert job["total_items"] == 3
    items = stack.client.get(f"/api/v1/jobs/{job['job_id']}/items").json()["items"]
    by_name = {item["raw_text"].split(" ")[0]: item for item in items}  # Ноутбук, Принтер, Бумага

    response = _export(stack, job["job_id"])
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    assert response.headers["x-unreviewed-count"] == "3"

    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    cells = [[c.value for c in row] for row in sheet.iter_rows()]
    assert cells[1][4:] == ADDED  # новые колонки — в строке заголовка, а не в строке названия
    for row in cells[2:]:
        name = (row[0] or "").split(" ")[0]
        expected = by_name[name]["final"]["ktru_code"] if name in by_name else None
        assert (row[4] or None) == expected, row  # пустая строка и «Итого» остаются без кода


def test_csv_and_json(stack: Stack) -> None:
    # «=…» openpyxl записал бы формулой без значения; «@» в csv так же опасен, а в xlsx остаётся текстом.
    risky = ["@SUM(A1) позиция с формулой", "", 1, "шт"]
    job = _job(stack, [HEADER, *ROWS, risky])

    response = _export(stack, job["job_id"], format="csv")
    assert response.headers["content-type"].startswith("text/csv")
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == HEADER + ADDED
    assert rows[4][0] == "'@SUM(A1) позиция с формулой"  # формула не выполнится в Excel

    data = json.loads(_export(stack, job["job_id"], format="json").content)
    assert [row[HEADER[0]] for row in data] == [r[0] for r in [*ROWS, risky]]
    assert all(row["ktru_code"] and row["source"] == "prediction" for row in data)
    assert all(len(str(row["confidence"]).split(".")[-1]) <= 4 for row in data)  # без хвоста float32

    _error(_export(stack, job["job_id"], format="docx"), 400, "invalid_request")


def test_unreviewed_items_and_job_status(stack: Stack) -> None:
    job = _job(stack, [HEADER, *ROWS, ["Позиция __fail__", "", 1, "шт"]])
    job_id = job["job_id"]
    error = _error(_export(stack, job_id, allow_unreviewed="false"), 409, "unreviewed_items")
    assert error["details"] == {"unreviewed": 3}

    stack.client.post(f"/api/v1/jobs/{job_id}/confirm-bulk", json={"min_confidence": 0, "skip_flagged": False})
    response = _export(stack, job_id, allow_unreviewed="false")
    assert (response.status_code, response.headers["x-unreviewed-count"]) == (200, "0")  # упавшая — не «непроверенная»
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    statuses = [row[7].value for row in sheet.iter_rows(min_row=2)]
    assert statuses == ["Подтверждено", "Подтверждено", "Подтверждено", "Ошибка"]

    queued = stack.create_job(stack.upload(xlsx_bytes([HEADER, *ROWS]))["file_id"])
    error = _error(_export(stack, queued["job_id"]), 409, "invalid_state")
    assert (error["message"], error["details"]) == ("Задача ещё обрабатывается.", {"status": "queued"})

    stack.parse_only(queued["job_id"])
    stack.client.post(f"/api/v1/jobs/{queued['job_id']}/cancel")
    cancelled = openpyxl.load_workbook(io.BytesIO(_export(stack, queued["job_id"]).content)).active
    assert {row[7].value for row in cancelled.iter_rows(min_row=2)} == {"Не обработана"}

    _error(_export(stack, str(uuid.uuid4())), 404, "not_found")
