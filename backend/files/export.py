"""Генерация файлов экспорта результатов классификации (FR-7.1, FR-7.2).

Поддерживаемые форматы:
- source: формат исходного файла (для PDF выдаётся XLSX).
- xlsx: таблица Excel со строковыми ячейками для защиты от исполнения формул.
- csv: текстовый файл с UTF-8 BOM, разделителем и экранированием формул апострофом.
- json: структурированный JSON-массив объектов.

Безопасность (Formula Injection):
- В CSV значения, начинающиеся со знаков '=', '+', '-', '@', экранируются ведущим апострофом.
- В XLSX все значения записываются ячейками строкового типа (cell.data_type = 's').
- Заголовок X-Unreviewed-Count содержит количество непроверенных позиций.
- При allow_unreviewed=False и наличии непроверенных строк генерируется 409 unreviewed_items.
"""

import csv
import io
import json
from copy import copy
from pathlib import Path
from typing import Any, Dict, List, Optional, Union
from urllib.parse import quote

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.files.clean import clean_cell
from backend.files.exceptions import UnreviewedItemsError, UnsupportedTypeError
from backend.files.schemas import ExportFormat, ExportItem, ExportResult

ADDED_COLUMNS: List[str] = [
    "Код КТРУ",
    "Наименование КТРУ",
    "Уверенность (%)",
    "Статус",
    "Источник",
    "Предупреждения",
]

STATUS_LABELS: Dict[str, str] = {
    "confirmed": "Подтверждено",
    "corrected": "Исправлено",
    "processed": "Не проверено",
    "needs_review": "Требует проверки",
    "failed": "Ошибка",
    "pending": "Не обработана",
}

SOURCE_LABELS: Dict[str, str] = {
    "user": "Пользователь",
    "prediction": "Модель",
}

FLAG_LABELS: Dict[str, str] = {
    "required_missing": "Отсутствуют обязательные характеристики",
    "archived_code": "Архивный код КТРУ",
    "low_extraction_confidence": "Низкая уверенность извлечения",
}

UNREVIEWED_STATUSES = {"processed", "needs_review"}


def escape_formula_for_csv(val: Any) -> str:
    """
    Экранирует значения для CSV: если ячейка начинается с '=', '+', '-', '@',
    добавляется ведущий апостроф `'` для предотвращения исполнения формул в Excel/Calc.
    """
    if val is None:
        return ""
    text = str(val)
    if text and text[0] in ("=", "+", "-", "@"):
        return f"'{text}"
    return text


def format_content_disposition(filename: str) -> str:
    """
    Формирует заголовок Content-Disposition по RFC 5987 / RFC 6266
    с корректной кодировкой кириллических имён файлов.
    Запасное имя filename="..." должно быть строго ASCII (latin-1) для Starlette/ASGI.
    """
    ascii_name = "".join(c if (c.isascii() and (c.isalnum() or c in "._-")) else "_" for c in filename) or "export"
    quoted_name = quote(filename.encode("utf-8"))
    return f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{quoted_name}'


def _format_flags(flags: List[str]) -> str:
    """Форматирует список предупреждений в читаемую строку."""
    if not flags:
        return ""
    return ", ".join(FLAG_LABELS.get(f, f) for f in flags)


def _normalize_item(item: Union[ExportItem, Dict[str, Any]]) -> ExportItem:
    """Приводит элемент к модели ExportItem."""
    if isinstance(item, ExportItem):
        return item
    if isinstance(item, dict):
        return ExportItem(**item)
    raise ValueError(f"Неподдерживаемый тип элемента экспорта: {type(item)}")


def _extract_added_values(item: ExportItem) -> List[str]:
    """Возвращает значения для 6 добавленных колонок по правилам OpenAPI."""
    if item.status == "failed":
        return ["", "", "", "Ошибка", "", _format_flags(item.flags)]

    if item.status == "pending":
        return ["", "", "", "Не обработана", "", ""]

    ktru_code = item.ktru_code or ""
    ktru_name = item.ktru_name or ""
    conf_str = f"{round(item.confidence * 100)}%" if item.confidence is not None else ""
    status_str = STATUS_LABELS.get(item.status, item.status)
    source_str = SOURCE_LABELS.get(item.source, item.source or "")
    flags_str = _format_flags(item.flags)

    return [ktru_code, ktru_name, conf_str, status_str, source_str, flags_str]


def export_csv(
    headers: List[str],
    items: List[ExportItem],
    delimiter: str = ";",
) -> bytes:
    """
    Экспортирует результат в CSV с экранированием формул и кодировкой UTF-8 BOM.
    """
    output = io.StringIO()
    writer = csv.writer(output, delimiter=delimiter, quoting=csv.QUOTE_MINIMAL)

    # Заголовок
    full_headers = headers + ADDED_COLUMNS
    writer.writerow([escape_formula_for_csv(h) for h in full_headers])

    # Строки данных
    for item in items:
        raw_cells = [escape_formula_for_csv(c) for c in item.raw_row]
        added_cells = [escape_formula_for_csv(c) for c in _extract_added_values(item)]
        writer.writerow(raw_cells + added_cells)

    # UTF-8 с BOM для корректного открытия в Excel на Windows
    return output.getvalue().encode("utf-8-sig")


def export_xlsx(
    headers: List[str],
    items: List[Union[ExportItem, Dict[str, Any]]],
    sheet_name: Optional[str] = None,
    source_path: Optional[Union[str, Path]] = None,
    header_row: Optional[int] = None,
) -> bytes:
    """
    Экспортирует результат в XLSX. Все ячейки сохраняются строкового типа ('s'),
    что полностью исключает исполнение формул без визуального добавления апострофа.
    """
    normalized_items = [_normalize_item(i) for i in items]

    # Если передан исходный XLSX-файл, пробуем дополнить его существующий лист
    if source_path:
        p = Path(source_path)
        if p.exists() and p.suffix.lower() == ".xlsx":
            try:
                wb = openpyxl.load_workbook(str(p))
                ws = wb[sheet_name] if (sheet_name and sheet_name in wb.sheetnames) else wb.active

                # Определяем строку заголовка в Excel (1-based)
                if header_row is not None:
                    effective_header_row = max(0, header_row)
                else:
                    effective_header_row = 0
                    if headers:
                        clean_target_headers = [clean_cell(h) for h in headers]
                        for r in range(1, min(ws.max_row + 1, 50)):
                            row_vals = [clean_cell(ws.cell(row=r, column=c).value) for c in range(1, len(headers) + 1)]
                            if row_vals == clean_target_headers:
                                effective_header_row = r - 1
                                break

                header_excel_row = effective_header_row + 1

                # Дописываем колонки справа
                start_col = ws.max_column + 1
                ref_header_cell = ws.cell(row=header_excel_row, column=start_col - 1) if start_col > 1 else None

                for c_idx, h in enumerate(ADDED_COLUMNS):
                    cell = ws.cell(row=header_excel_row, column=start_col + c_idx, value=str(h))
                    cell.data_type = "s"
                    if ref_header_cell and ref_header_cell.has_style:
                        try:
                            if ref_header_cell.font:
                                cell.font = copy(ref_header_cell.font)
                                cell.font.bold = True
                            else:
                                cell.font = Font(bold=True)
                            if ref_header_cell.fill:
                                cell.fill = copy(ref_header_cell.fill)
                            if ref_header_cell.alignment:
                                cell.alignment = copy(ref_header_cell.alignment)
                            if ref_header_cell.border:
                                cell.border = copy(ref_header_cell.border)
                        except Exception:
                            cell.font = Font(bold=True)
                    else:
                        cell.font = Font(bold=True)

                for item in normalized_items:
                    target_row = header_excel_row + 1 + item.row_index
                    ref_data_cell = ws.cell(row=target_row, column=start_col - 1) if start_col > 1 else None
                    added_vals = _extract_added_values(item)
                    for c_idx, val in enumerate(added_vals):
                        cell = ws.cell(row=target_row, column=start_col + c_idx, value=str(val))
                        cell.data_type = "s"
                        if ref_data_cell and ref_data_cell.has_style:
                            try:
                                if ref_data_cell.font:
                                    cell.font = copy(ref_data_cell.font)
                                if ref_data_cell.alignment:
                                    cell.alignment = copy(ref_data_cell.alignment)
                                if ref_data_cell.border:
                                    cell.border = copy(ref_data_cell.border)
                            except Exception:
                                pass

                # Автоподбор ширины для добавленных колонок
                for c_idx in range(len(ADDED_COLUMNS)):
                    col_idx = start_col + c_idx
                    col_letter = get_column_letter(col_idx)
                    max_len = 0
                    for r in range(1, ws.max_row + 1):
                        v = ws.cell(row=r, column=col_idx).value
                        if v:
                            max_len = max(max_len, len(str(v)))
                    ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

                buf = io.BytesIO()
                wb.save(buf)
                return buf.getvalue()
            except Exception:
                # Если исходный файл не открылся openpyxl, создаём чистую таблицу
                pass

    # Создаём новую книгу
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name or "Результаты"

    # Заголовки
    full_headers = headers + ADDED_COLUMNS
    ws.append(full_headers)

    header_font = Font(bold=True, color="000000")
    header_fill = PatternFill(start_color="F0F2F5", end_color="F0F2F5", fill_type="solid")

    for col_idx in range(1, len(full_headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.data_type = "s"

    # Строки данных
    for r_idx, item in enumerate(normalized_items, start=2):
        row_vals = item.raw_row + _extract_added_values(item)
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=str(val) if val is not None else "")
            cell.data_type = "s"

    # Автоподбор ширины колонок
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_json(
    headers: List[str],
    items: List[ExportItem],
) -> bytes:
    """
    Экспортирует результат в JSON (массив объектов с исходными полями и КТРУ).
    """
    rows_data: List[Dict[str, Any]] = []

    for item in items:
        row_dict: Dict[str, Any] = {}
        for idx, val in enumerate(item.raw_row):
            col_key = headers[idx] if idx < len(headers) else f"column_{idx}"
            row_dict[col_key] = val

        if item.status == "failed":
            row_dict["ktru_code"] = None
            row_dict["ktru_name"] = None
            row_dict["confidence"] = None
            row_dict["status"] = "failed"
            row_dict["status_label"] = "Ошибка"
            row_dict["source"] = None
            row_dict["flags"] = item.flags
        elif item.status == "pending":
            row_dict["ktru_code"] = None
            row_dict["ktru_name"] = None
            row_dict["confidence"] = None
            row_dict["status"] = "pending"
            row_dict["status_label"] = "Не обработана"
            row_dict["source"] = None
            row_dict["flags"] = []
        else:
            row_dict["ktru_code"] = item.ktru_code
            row_dict["ktru_name"] = item.ktru_name
            row_dict["confidence"] = item.confidence
            row_dict["status"] = item.status
            row_dict["status_label"] = STATUS_LABELS.get(item.status, item.status)
            row_dict["source"] = item.source
            row_dict["flags"] = item.flags

        rows_data.append(row_dict)

    return json.dumps(rows_data, ensure_ascii=False, indent=2).encode("utf-8")


def export_job_results(
    items: List[Union[ExportItem, Dict[str, Any]]],
    headers: Optional[List[str]] = None,
    export_format: Union[str, ExportFormat] = ExportFormat.SOURCE,
    original_filename: str = "specification.xlsx",
    source_path: Optional[Union[str, Path]] = None,
    sheet_name: Optional[str] = None,
    allow_unreviewed: bool = True,
    csv_delimiter: str = ";",
    header_row: Optional[int] = None,
) -> ExportResult:
    """
    Универсальная функция выгрузки результатов задачи обработки спецификации (FR-7.1, FR-7.2).

    :param items: Список позиций (ExportItem или словари)
    :param headers: Заголовки исходных колонок
    :param export_format: Формат: 'source', 'xlsx', 'csv', 'json'
    :param original_filename: Имя исходного файла спецификации
    :param source_path: Путь к исходному файлу на диске (если есть)
    :param sheet_name: Название листа спецификации
    :param allow_unreviewed: Разрешить ли выгрузку при наличии непроверенных позиций
    :param csv_delimiter: Разделитель для формата CSV (по умолчанию ';')
    :param header_row: Номер строки заголовка в исходном файле (с 0)
    :return: ExportResult с байтами, media_type, именем файла и заголовками
    """
    normalized_items = [_normalize_item(i) for i in items]

    # Считаем количество непроверенных позиций (processed и needs_review)
    unreviewed_count = sum(1 for i in normalized_items if i.status in UNREVIEWED_STATUSES)

    if not allow_unreviewed and unreviewed_count > 0:
        raise UnreviewedItemsError(
            f"Осталось {unreviewed_count} непроверенных позиций.",
            details={"unreviewed": unreviewed_count},
        )

    # Определяем результирующий формат
    fmt_str = export_format.value if isinstance(export_format, ExportFormat) else str(export_format).lower()
    orig_path = Path(original_filename)
    stem = orig_path.stem or "export"
    orig_ext = orig_path.suffix.lower()

    if fmt_str == ExportFormat.SOURCE.value:
        if orig_ext == ".csv":
            target_format = ExportFormat.CSV
        elif orig_ext in (".xlsx", ".xls"):
            target_format = ExportFormat.XLSX
        elif orig_ext == ".pdf":
            # По контракту: для PDF исходный формат выгружается как XLSX
            target_format = ExportFormat.XLSX
        else:
            target_format = ExportFormat.XLSX
    elif fmt_str == ExportFormat.XLSX.value:
        target_format = ExportFormat.XLSX
    elif fmt_str == ExportFormat.CSV.value:
        target_format = ExportFormat.CSV
    elif fmt_str == ExportFormat.JSON.value:
        target_format = ExportFormat.JSON
    else:
        raise UnsupportedTypeError(f"Неподдерживаемый формат экспорта: {export_format}")

    # Восстанавливаем заголовки по первой строке данных, если не переданы
    if headers is None:
        if normalized_items and normalized_items[0].raw_row:
            headers = [f"Колонка {i + 1}" for i in range(len(normalized_items[0].raw_row))]
        else:
            headers = []

    # Генерация содержимого
    if target_format == ExportFormat.XLSX:
        out_filename = f"{stem}_ktru.xlsx"
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        content = export_xlsx(
            headers,
            normalized_items,
            sheet_name=sheet_name,
            source_path=source_path,
            header_row=header_row,
        )
    elif target_format == ExportFormat.CSV:
        out_filename = f"{stem}_ktru.csv"
        media_type = "text/csv; charset=utf-8"
        content = export_csv(headers, normalized_items, delimiter=csv_delimiter)
    elif target_format == ExportFormat.JSON:
        out_filename = f"{stem}_ktru.json"
        media_type = "application/json; charset=utf-8"
        content = export_json(headers, normalized_items)

    content_disposition = format_content_disposition(out_filename)

    return ExportResult(
        content=content,
        media_type=media_type,
        filename=out_filename,
        unreviewed_count=unreviewed_count,
        content_disposition=content_disposition,
    )
