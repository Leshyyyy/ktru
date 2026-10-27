"""Тесты модуля экспорта результатов (FR-7.1, FR-7.2)."""

import io
import json
import openpyxl
import pytest

from backend.files.exceptions import UnreviewedItemsError
from backend.files.export import (
    ADDED_COLUMNS,
    escape_formula_for_csv,
    export_csv,
    export_job_results,
    export_json,
    export_xlsx,
    format_content_disposition,
)
from backend.files.parser import parse_rows
from backend.files.schemas import (
    ColumnMappingItem,
    ColumnRole,
    ExportFormat,
    ExportItem,
)


@pytest.fixture
def sample_export_items():
    return [
        ExportItem(
            row_index=0,
            raw_row=["Ноутбук Lenovo IdeaPad 3", "12", "шт"],
            ktru_code="26.20.11.110-00000001",
            ktru_name="Компьютер портативный (ноутбук)",
            confidence=0.942,
            status="confirmed",
            source="user",
            flags=[],
        ),
        ExportItem(
            row_index=1,
            raw_row=["Принтер Xerox B230", "3", "шт"],
            ktru_code="26.20.16.120-00000002",
            ktru_name="Принтер лазерный",
            confidence=0.875,
            status="processed",
            source="prediction",
            flags=["low_extraction_confidence"],
        ),
        ExportItem(
            row_index=2,
            raw_row=["Непонятная позиция", "1", "компл"],
            ktru_code=None,
            ktru_name=None,
            confidence=None,
            status="failed",
            source=None,
            flags=[],
            error="Не удалось извлечь атрибуты",
        ),
        ExportItem(
            row_index=3,
            raw_row=["Отменённая строка", "5", "шт"],
            ktru_code=None,
            ktru_name=None,
            confidence=None,
            status="pending",
            source=None,
            flags=[],
        ),
    ]


@pytest.fixture
def sample_headers():
    return ["Наименование", "Кол-во", "Ед. изм."]


def test_escape_formula_for_csv():
    """Проверка экранирования CSV-инъекций (=, +, -, @)."""
    assert escape_formula_for_csv("=1+1") == "'=1+1"
    assert escape_formula_for_csv("+SUM(A1)") == "'+SUM(A1)"
    assert escape_formula_for_csv("-100") == "'-100"
    assert escape_formula_for_csv("@calc") == "'@calc"
    assert escape_formula_for_csv("Обычный текст") == "Обычный текст"
    assert escape_formula_for_csv(123) == "123"
    assert escape_formula_for_csv(None) == ""


def test_format_content_disposition():
    """Проверка формирования заголовка Content-Disposition по RFC 5987 / RFC 6266."""
    header = format_content_disposition("спецификация_2026.xlsx")
    assert 'attachment; filename=' in header
    assert "filename*=UTF-8''" in header
    assert "%D1%81%D0%BF%D0%B5%D1%86%D0%B8%D1%84%D0%B8%D0%BA%D0%B0%D1%86%D0%B8%D1%8F" in header

    # Заголовок должен быть строго ASCII и успешно кодироваться в latin-1 (требование Starlette/HTTP)
    header_bytes = header.encode("latin-1")
    assert isinstance(header_bytes, bytes)
    assert header.isascii()

    # Запасное имя filename="..." должно состоять только из ASCII (без русских букв)
    import re
    match = re.search(r'filename="([^"]+)"', header)
    assert match is not None
    fallback_name = match.group(1)
    assert fallback_name.isascii()
    assert fallback_name == "_____________2026.xlsx"


def test_export_csv_structure_and_bom(sample_headers, sample_export_items):
    """Проверка структуры CSV, экранирования формул и UTF-8 BOM."""
    csv_bytes = export_csv(sample_headers, sample_export_items, delimiter=";")
    assert csv_bytes.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM

    text = csv_bytes.decode("utf-8-sig")
    lines = text.strip().split("\r\n")
    if len(lines) == 1:
        lines = text.strip().split("\n")

    # Проверка строки заголовков
    header_line = lines[0].split(";")
    expected_headers = sample_headers + ADDED_COLUMNS
    assert header_line == expected_headers

    # Строка 1: confirmed, user
    row1 = lines[1].split(";")
    assert row1[0] == "Ноутбук Lenovo IdeaPad 3"
    assert row1[3] == "26.20.11.110-00000001"
    assert row1[5] == "94%"
    assert row1[6] == "Подтверждено"
    assert row1[7] == "Пользователь"

    # Строка 2: processed, prediction, flag
    row2 = lines[2].split(";")
    assert row2[5] == "88%"
    assert row2[6] == "Не проверено"
    assert row2[7] == "Модель"
    assert "Низкая уверенность извлечения" in row2[8]

    # Строка 3: failed
    row3 = lines[3].split(";")
    assert row3[3] == ""  # ktru_code пустой
    assert row3[4] == ""  # ktru_name пустое
    assert row3[6] == "Ошибка"

    # Строка 4: pending
    row4 = lines[4].split(";")
    assert row4[6] == "Не обработана"


def test_export_xlsx_string_cells(sample_headers, sample_export_items):
    """Проверка XLSX: открывается, содержит данные и все ячейки строго строкового типа (data_type='s')."""
    # Добавляем потенциально опасную строку с формулой
    items_with_formula = sample_export_items + [
        ExportItem(
            row_index=4,
            raw_row=["=cmd|' /C calc'!A0", "1", "шт"],
            ktru_code="26.20.11.110-00000001",
            ktru_name="Тест",
            confidence=0.9,
            status="confirmed",
        )
    ]

    xlsx_bytes = export_xlsx(sample_headers, items_with_formula)
    assert len(xlsx_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    assert ws.title == "Результаты"

    # Проверяем, что все ячейки строки имеют строковый тип (не 'f' формула)
    for row in ws.iter_rows():
        for cell in row:
            assert cell.data_type in ("s", "inlineStr", "str"), f"Ячейка {cell.coordinate} не должна быть формулой, текущий тип: {cell.data_type}"

    # Проверяем опасную формулу — она должна сохраниться как обычный текст со строковым типом
    formula_cell = ws.cell(row=6, column=1)
    assert formula_cell.data_type == "s"
    assert formula_cell.value == "=cmd|' /C calc'!A0"


def test_export_json_format(sample_headers, sample_export_items):
    """Проверка выгрузки в JSON."""
    json_bytes = export_json(sample_headers, sample_export_items)
    data = json.loads(json_bytes.decode("utf-8"))

    assert len(data) == 4
    assert data[0]["Наименование"] == "Ноутбук Lenovo IdeaPad 3"
    assert data[0]["ktru_code"] == "26.20.11.110-00000001"
    assert data[0]["status"] == "confirmed"
    assert data[0]["status_label"] == "Подтверждено"
    assert data[0]["source"] == "user"

    # failed позиция
    assert data[2]["ktru_code"] is None
    assert data[2]["status"] == "failed"
    assert data[2]["status_label"] == "Ошибка"

    # pending позиция
    assert data[3]["status"] == "pending"
    assert data[3]["status_label"] == "Не обработана"


def test_export_job_results_unreviewed_count_and_header(sample_headers, sample_export_items):
    """Проверка подсчёта X-Unreviewed-Count и заголовка."""
    # В sample_export_items одна позиция 'processed'
    res = export_job_results(
        sample_export_items,
        sample_headers,
        export_format=ExportFormat.XLSX,
        original_filename="spec.xlsx",
        allow_unreviewed=True,
    )

    assert res.unreviewed_count == 1
    assert res.filename == "spec_ktru.xlsx"
    assert res.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "filename*=UTF-8''spec_ktru.xlsx" in res.content_disposition


def test_export_job_results_allow_unreviewed_false(sample_headers, sample_export_items):
    """При allow_unreviewed=False и наличии непроверенных строк генерируется UnreviewedItemsError (409)."""
    with pytest.raises(UnreviewedItemsError) as exc_info:
        export_job_results(
            sample_export_items,
            sample_headers,
            export_format=ExportFormat.CSV,
            allow_unreviewed=False,
        )

    assert exc_info.value.code == "unreviewed_items"
    assert exc_info.value.details == {"unreviewed": 1}
    assert "Осталось 1 непроверенных" in str(exc_info.value)


def test_export_job_results_format_source(sample_headers, sample_export_items):
    """Проверка автоматического выбора формата при format='source'."""
    # Для CSV -> CSV
    res_csv = export_job_results(
        sample_export_items,
        sample_headers,
        export_format=ExportFormat.SOURCE,
        original_filename="data.csv",
    )
    assert res_csv.filename.endswith(".csv")
    assert "text/csv" in res_csv.media_type

    # Для XLSX -> XLSX
    res_xlsx = export_job_results(
        sample_export_items,
        sample_headers,
        export_format=ExportFormat.SOURCE,
        original_filename="data.xlsx",
    )
    assert res_xlsx.filename.endswith(".xlsx")
    assert "spreadsheetml" in res_xlsx.media_type

    # Для PDF -> XLSX (по контракту)
    res_pdf = export_job_results(
        sample_export_items,
        sample_headers,
        export_format=ExportFormat.SOURCE,
        original_filename="document.pdf",
    )
    assert res_pdf.filename.endswith(".xlsx")
    assert "spreadsheetml" in res_pdf.media_type


def test_export_xlsx_source_file_with_header_row_and_gaps(tmp_path):
    """
    Тест краевого случая: экспорт в исходный XLSX файл с шапкой над таблицей (header_row > 0)
    и пустыми строками между позициями.
    Проверяет:
    1. Колонки КТРУ дописываются в строку заголовков таблицы (header_row), а не в шапку над ней.
    2. Коды КТРУ попадают в строки соответствующих позиций по item.row_index.
    3. Пустые строки и шапка над таблицей не затрагиваются.
    """
    xlsx_path = tmp_path / "spec_test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Спецификация"

    # Строка 1: Шапка над таблицей
    ws.cell(row=1, column=1, value="Спецификация №1")
    # Строка 2: Заголовок таблицы (header_row = 1 в 0-based индексации)
    ws.cell(row=2, column=1, value="Наименование")
    ws.cell(row=2, column=2, value="Кол-во")
    # Строка 3: Позиция 1 (row_index = 0 после заголовка)
    ws.cell(row=3, column=1, value="Перчатки")
    ws.cell(row=3, column=2, value="10")
    # Строка 4: Пустая строка (row_index = 1 после заголовка, оставляем пустой)
    # Строка 5: Позиция 2 (row_index = 2 после заголовка)
    ws.cell(row=5, column=1, value="Бумага А4")
    ws.cell(row=5, column=2, value="5")

    wb.save(xlsx_path)

    items = [
        ExportItem(
            row_index=0,
            raw_row=["Перчатки", "10"],
            ktru_code="КОД-Перчатки",
            ktru_name="Перчатки медицинские",
            confidence=0.95,
            status="confirmed",
        ),
        ExportItem(
            row_index=2,
            raw_row=["Бумага А4", "5"],
            ktru_code="КОД-Бумага А4",
            ktru_name="Бумага офисная",
            confidence=0.88,
            status="processed",
        ),
    ]
    headers = ["Наименование", "Кол-во"]

    # 1. Проверяем явную передачу header_row=1
    res_bytes = export_xlsx(
        headers=headers,
        items=items,
        sheet_name="Спецификация",
        source_path=xlsx_path,
        header_row=1,
    )

    wb_out = openpyxl.load_workbook(io.BytesIO(res_bytes))
    ws_out = wb_out["Спецификация"]

    # Строка 1: "Спецификация №1", колонка 3 пустая (код не попал в шапку)
    assert ws_out.cell(row=1, column=1).value == "Спецификация №1"
    assert ws_out.cell(row=1, column=3).value is None

    # Строка 2: Заголовок таблицы -> здесь должны появиться заголовки добавленных колонок
    assert ws_out.cell(row=2, column=1).value == "Наименование"
    assert ws_out.cell(row=2, column=2).value == "Кол-во"
    assert ws_out.cell(row=2, column=3).value == "Код КТРУ"

    # Строка 3: "Перчатки" -> код Перчаток
    assert ws_out.cell(row=3, column=1).value == "Перчатки"
    assert ws_out.cell(row=3, column=2).value == "10"
    assert ws_out.cell(row=3, column=3).value == "КОД-Перчатки"

    # Строка 4: Пустая строка остаётся пустой
    assert ws_out.cell(row=4, column=1).value is None
    assert ws_out.cell(row=4, column=3).value is None

    # Строка 5: "Бумага А4" -> код Бумаги
    assert ws_out.cell(row=5, column=1).value == "Бумага А4"
    assert ws_out.cell(row=5, column=2).value == "5"
    assert ws_out.cell(row=5, column=3).value == "КОД-Бумага А4"

    # 2. Проверяем автоопределение header_row по headers при header_row=None
    res_auto = export_xlsx(
        headers=headers,
        items=items,
        sheet_name="Спецификация",
        source_path=xlsx_path,
        header_row=None,
    )
    wb_auto = openpyxl.load_workbook(io.BytesIO(res_auto))
    ws_auto = wb_auto["Спецификация"]
    assert ws_auto.cell(row=1, column=3).value is None
    assert ws_auto.cell(row=2, column=3).value == "Код КТРУ"
    assert ws_auto.cell(row=3, column=3).value == "КОД-Перчатки"
    assert ws_auto.cell(row=4, column=3).value is None
    assert ws_auto.cell(row=5, column=3).value == "КОД-Бумага А4"

    # 3. Проверяем через export_job_results с форматом SOURCE
    res_job = export_job_results(
        items=items,
        headers=headers,
        export_format=ExportFormat.SOURCE,
        original_filename="spec_test.xlsx",
        source_path=xlsx_path,
        sheet_name="Спецификация",
        header_row=1,
    )
    wb_job = openpyxl.load_workbook(io.BytesIO(res_job.content))
    ws_job = wb_job["Спецификация"]
    assert ws_job.cell(row=1, column=3).value is None
    assert ws_job.cell(row=2, column=3).value == "Код КТРУ"
    assert ws_job.cell(row=3, column=3).value == "КОД-Перчатки"
    assert ws_job.cell(row=4, column=3).value is None
    assert ws_job.cell(row=5, column=3).value == "КОД-Бумага А4"


def test_parse_and_export_roundtrip_xlsx(tmp_path):
    """
    Сквозной тест разбора и экспорта (parse_rows -> ExportItem -> export_xlsx).
    Имитирует полный цикл обработки файла спецификации с шапкой и пустыми строками.
    """
    xlsx_path = tmp_path / "roundtrip.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Спецификация"

    ws.append(["Спецификация №1", ""])  # Строка 1: заголовок документа
    ws.append(["Наименование", "Кол-во"])  # Строка 2: заголовок таблицы
    ws.append(["Перчатки", "10"])  # Строка 3: позиция
    ws.append(["", ""])  # Строка 4: пустая строка
    ws.append(["Бумага А4", "5"])  # Строка 5: позиция
    wb.save(xlsx_path)

    mapping = [
        ColumnMappingItem(index=0, role=ColumnRole.NAME),
        ColumnMappingItem(index=1, role=ColumnRole.QUANTITY),
    ]

    # 1. Разбираем файл
    parsed_rows = parse_rows(xlsx_path, sheet_name="Спецификация", header_row=1, column_mapping=mapping)
    assert len(parsed_rows) == 2
    assert parsed_rows[0].row_index == 0
    assert parsed_rows[0].raw_text == "Перчатки"
    assert parsed_rows[1].row_index == 2
    assert parsed_rows[1].raw_text == "Бумага А4"

    # 2. Формируем результаты классификации
    export_items = [
        ExportItem(
            row_index=parsed_rows[0].row_index,
            raw_row=parsed_rows[0].raw_row,
            ktru_code="КОД-Перчатки",
            ktru_name="Перчатки медицинские",
            confidence=0.99,
            status="confirmed",
        ),
        ExportItem(
            row_index=parsed_rows[1].row_index,
            raw_row=parsed_rows[1].raw_row,
            ktru_code="КОД-Бумага А4",
            ktru_name="Бумага для печати",
            confidence=0.91,
            status="processed",
        ),
    ]

    # 3. Экспортируем в исходный файл
    out_bytes = export_xlsx(
        headers=["Наименование", "Кол-во"],
        items=export_items,
        sheet_name="Спецификация",
        source_path=xlsx_path,
        header_row=1,
    )

    wb_res = openpyxl.load_workbook(io.BytesIO(out_bytes))
    ws_res = wb_res["Спецификация"]

    # Проверяем строки:
    # 1. Шапка не задета
    assert [ws_res.cell(row=1, column=c).value for c in range(1, 4)] == ["Спецификация №1", None, None]
    # 2. Заголовок таблицы расширен добавленными колонками
    assert ws_res.cell(row=2, column=1).value == "Наименование"
    assert ws_res.cell(row=2, column=2).value == "Кол-во"
    assert ws_res.cell(row=2, column=3).value == "Код КТРУ"
    # 3. Перчатки получили свой код
    assert ws_res.cell(row=3, column=1).value == "Перчатки"
    assert ws_res.cell(row=3, column=2).value == "10"
    assert ws_res.cell(row=3, column=3).value == "КОД-Перчатки"
    # 4. Пустая строка осталась пустой
    assert [ws_res.cell(row=4, column=c).value for c in range(1, 4)] == [None, None, None]
    # 5. Бумага А4 получила свой код
    assert ws_res.cell(row=5, column=1).value == "Бумага А4"
    assert ws_res.cell(row=5, column=2).value == "5"
    assert ws_res.cell(row=5, column=3).value == "КОД-Бумага А4"
