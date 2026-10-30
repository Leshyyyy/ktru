"""End-to-End тест полного конвейера обработки спецификации (FR-1, FR-2, FR-6, FR-7).

Сквозной сценарий:
1. Загрузка файла (read_file):
   - Извлечение структуры, листов, превью строк.
   - Хеширование: вычисление sha256 (дубликаты файлов) и хеша заголовков (layout_hint).
   - Автоматическое определение ролей колонок (suggest_roles).
2. Разбор файла (parse_rows):
   - Проверка обязательной роли 'name' и формирование column_mapping.
   - Склейка raw_text (name + description + characteristics) и сохранение raw_row.
   - Фильтрация пустых строк и строк итогов с сохранением row_index.
3. Классификация ML (StubPipeline):
   - Пакетная обработка позиций пачками по 50 штук (имитация Celery-воркера).
   - Получение предсказаний (топ-3 КТРУ, уверенность, извлечение атрибутов, зоны).
   - Обработка сбойных позиций (маркер __fail__ -> статус failed).
4. Ручная правка и подтверждение:
   - Подтверждение кода для одной позиции (confirm -> confirmed, источник prediction).
   - Ручная корректировка кода для другой позиции (decision -> corrected, источник user).
   - Формирование итогового набора данных по логике представления v_item_final.
5. Экспорт результатов (export_job_results):
   - Экспорт в XLSX, CSV, JSON и SOURCE.
   - Проверка наличия исходных колонок + 6 добавленных колонок.
   - Проверка, что пользовательские правки и подтверждения попали в выгрузку.
   - Проверка защиты от Formula Injection (экранирование '=', '+', '-', '@').
   - Проверка заголовка X-Unreviewed-Count и блокировки allow_unreviewed=False.
   - Валидация сгенерированного файла XLSX через Calamine и OpenPyXL.
"""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import List
from uuid import uuid4

import openpyxl
import pytest
from python_calamine import CalamineWorkbook

from backend.files import (
    ColumnMappingItem,
    ColumnRole,
    ExportFormat,
    ExportItem,
    compute_file_sha256,
    compute_headers_hash,
    export_job_results,
    parse_rows,
    read_file,
    suggest_roles,
)
from backend.files.exceptions import UnreviewedItemsError
from contracts.ml_contracts import RawItem
from contracts.stub_pipeline import StubPipeline


@pytest.fixture(scope="session")
def fixtures_dir() -> Path:
    """Возвращает путь к директории с тестовыми спецификациями."""
    p = Path(__file__).resolve().parent.parent / "fixtures"
    if not (p / "specification_100.xlsx").exists():
        from tests.generate_test_specs import generate_all_test_fixtures

        generate_all_test_fixtures(p)
    return p


@pytest.fixture
def spec_100_xlsx(fixtures_dir: Path) -> Path:
    return fixtures_dir / "specification_100.xlsx"


@pytest.fixture
def spec_100_csv(fixtures_dir: Path) -> Path:
    return fixtures_dir / "specification_100.csv"


@pytest.fixture
def spec_1000_xlsx(fixtures_dir: Path) -> Path:
    return fixtures_dir / "specification_1000.xlsx"



def test_e2e_pipeline_xlsx_full_flow(spec_100_xlsx: Path, tmp_path: Path):
    """
    Полный сквозной тест для спецификации в формате XLSX:
    Загрузка -> Определение ролей -> Разбор строк -> ML-заглушка -> Правка -> Экспорт.
    """
    # -------------------------------------------------------------------------
    # ШАГ 1: Загрузка и предварительный просмотр (POST /files)
    # -------------------------------------------------------------------------
    file_info = read_file(spec_100_xlsx)
    assert file_info.filename == "specification_100.xlsx"
    assert file_info.size_bytes > 0
    assert len(file_info.sheets) == 1

    sheet = file_info.sheets[0]
    assert sheet.name == "Спецификация"
    assert sheet.row_count == 100
    assert len(sheet.columns) == 6
    assert len(sheet.preview) > 0

    # Проверка вычисления хешей для duplicate_of_job_id и layout_hint
    sha256 = compute_file_sha256(spec_100_xlsx)
    assert len(sha256) == 64

    headers = [c.header for c in sheet.columns]
    headers_hash = compute_headers_hash(headers)
    assert len(headers_hash) == 64

    # Проверка автосопоставления ролей (уже заполнены в sheet.columns)
    role_map = {c.index: c.suggested_role for c in sheet.columns}

    # Колонка 0 ("№ п/п") -> ignore
    # Колонка 1 ("Наименование товара...") -> name
    # Колонка 2 ("Описание...") -> description
    # Колонка 3 ("Характеристики") -> characteristics
    # Колонка 4 ("Количество") -> quantity
    # Колонка 5 ("Единица измерения") -> unit
    assert role_map[0] == ColumnRole.IGNORE
    assert role_map[1] == ColumnRole.NAME
    assert role_map[2] == ColumnRole.DESCRIPTION
    assert role_map[3] == ColumnRole.CHARACTERISTICS
    assert role_map[4] == ColumnRole.QUANTITY
    assert role_map[5] == ColumnRole.UNIT

    mapping = [ColumnMappingItem(index=c.index, role=c.suggested_role) for c in sheet.columns]

    # -------------------------------------------------------------------------
    # ШАГ 2: Разбор строк для создания задачи (POST /jobs & Celery parse_file)
    # -------------------------------------------------------------------------
    parsed_rows = parse_rows(
        path=spec_100_xlsx,
        sheet_name=sheet.name,
        header_row=0,
        column_mapping=mapping,
    )

    assert len(parsed_rows) == 100
    for idx, row in enumerate(parsed_rows):
        assert row.row_index == idx
        assert len(row.raw_row) == 6
        assert len(row.raw_text) > 0

    # Проверка корректной склейки name + description + characteristics
    first_row = parsed_rows[0]
    assert "Ноутбук Lenovo IdeaPad 3 15IAU7" in first_row.raw_text
    assert "Портативный компьютер для офисных задач" in first_row.raw_text
    assert "Экран 15.6" in first_row.raw_text

    # -------------------------------------------------------------------------
    # ШАГ 3: Классификация ML через StubPipeline (пачками по 50 позиций)
    # -------------------------------------------------------------------------
    pipeline = StubPipeline()
    assert pipeline.is_ready() == {"models": True, "catalog": True}

    chunk_size = 50
    classification_results = []

    for i in range(0, len(parsed_rows), chunk_size):
        chunk = parsed_rows[i : i + chunk_size]
        raw_items = [
            RawItem(
                item_id=str(uuid4()),
                text=r.raw_text,
                columns={
                    "name": r.raw_row[1],
                    "description": r.raw_row[2],
                    "characteristics": r.raw_row[3],
                    "quantity": r.raw_row[4],
                    "unit": r.raw_row[5],
                },
            )
            for r in chunk
        ]
        results = pipeline.classify_batch(raw_items, catalog_version_id="2026-10-25")
        assert len(results) == len(chunk)
        classification_results.extend(results)

    assert len(classification_results) == 100

    # Проверка, что контрольная строка с __fail__ (строка с индексом 7) вернула ошибку
    failed_res = classification_results[7]
    assert failed_res.error is not None
    assert len(failed_res.predictions) == 0

    # Проверка нормальных результатов классификации
    normal_res = classification_results[0]
    assert normal_res.error is None
    assert len(normal_res.predictions) == 3
    assert normal_res.predictions[0].ktru_code.startswith("99.99.")
    assert normal_res.predictions[0].confidence > 0.0
    assert normal_res.extracted.brand == "Lenovo"

    # -------------------------------------------------------------------------
    # ШАГ 4: Имитация решений оператора и формирование v_item_final
    # -------------------------------------------------------------------------
    export_items: List[ExportItem] = []

    for idx, (p_row, c_res) in enumerate(zip(parsed_rows, classification_results, strict=False)):
        item_id = uuid4()

        if c_res.error:
            export_items.append(
                ExportItem(
                    item_id=item_id,
                    row_index=p_row.row_index,
                    raw_row=p_row.raw_row,
                    status="failed",
                    error=c_res.error,
                    flags=[],
                )
            )
            continue

        top_pred = c_res.predictions[0]

        if idx == 0:
            # Пользователь подтверждает первое предсказание (POST /items/{id}/confirm)
            export_items.append(
                ExportItem(
                    item_id=item_id,
                    row_index=p_row.row_index,
                    raw_row=p_row.raw_row,
                    ktru_code=top_pred.ktru_code,
                    ktru_name=top_pred.name,
                    confidence=top_pred.confidence,
                    status="confirmed",
                    source="prediction",
                    flags=[],
                )
            )
        elif idx == 1:
            # Пользователь вручную исправляет код КТРУ (PUT /items/{id}/decision)
            manual_code = "99.99.01.000-00000002"
            manual_name = "Ноутбук игровой кастомный"
            export_items.append(
                ExportItem(
                    item_id=item_id,
                    row_index=p_row.row_index,
                    raw_row=p_row.raw_row,
                    ktru_code=manual_code,
                    ktru_name=manual_name,
                    confidence=1.0,
                    status="corrected",
                    source="user",
                    flags=[],
                )
            )
        else:
            # Позиции без ручного вмешательства (processed или needs_review)
            status = "processed" if top_pred.confidence >= 0.85 else "needs_review"
            export_items.append(
                ExportItem(
                    item_id=item_id,
                    row_index=p_row.row_index,
                    raw_row=p_row.raw_row,
                    ktru_code=top_pred.ktru_code,
                    ktru_name=top_pred.name,
                    confidence=top_pred.confidence,
                    status=status,
                    source="prediction",
                    flags=[],
                )
            )

    # -------------------------------------------------------------------------
    # ШАГ 5: Экспорт результатов (GET /jobs/{id}/export)
    # -------------------------------------------------------------------------
    # 5.1 Проверка запрета выгрузки при наличии непроверенных строк (allow_unreviewed=False)
    with pytest.raises(UnreviewedItemsError) as exc_info:
        export_job_results(
            items=export_items,
            headers=headers,
            export_format=ExportFormat.XLSX,
            original_filename="specification_100.xlsx",
            source_path=spec_100_xlsx,
            sheet_name=sheet.name,
            allow_unreviewed=False,
            header_row=0,
        )
    assert exc_info.value.details["unreviewed"] > 0

    # 5.2 Успешный экспорт в XLSX с добавлением колонок к исходному файлу
    res_xlsx = export_job_results(
        items=export_items,
        headers=headers,
        export_format=ExportFormat.XLSX,
        original_filename="specification_100.xlsx",
        source_path=spec_100_xlsx,
        sheet_name=sheet.name,
        allow_unreviewed=True,
        header_row=0,
    )

    assert res_xlsx.filename == "specification_100_ktru.xlsx"
    assert res_xlsx.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert res_xlsx.unreviewed_count > 0
    assert 'filename="specification_100_ktru.xlsx"' in res_xlsx.content_disposition
    assert "filename*=UTF-8''" in res_xlsx.content_disposition

    # 5.3 Валидация сформированного файла XLSX через openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
    ws = wb["Спецификация"]

    # Проверка числа строк и колонок (6 исходных + 6 добавленных)
    assert ws.max_row == 101  # 1 заголовок + 100 позиций
    assert ws.max_column == 12

    # Заголовки добавленных колонок
    added_headers = [ws.cell(row=1, column=c).value for c in range(7, 13)]
    assert added_headers == [
        "Код КТРУ",
        "Наименование КТРУ",
        "Уверенность (%)",
        "Статус",
        "Источник",
        "Предупреждения",
    ]

    # Строка 0 (Excel row 2): Подтверждено, Модель
    row2_vals = [ws.cell(row=2, column=c).value for c in range(7, 13)]
    assert row2_vals[3] == "Подтверждено"
    assert row2_vals[4] == "Модель"
    assert row2_vals[0].startswith("99.99.")

    # Строка 1 (Excel row 3): Исправлено, Пользователь
    row3_vals = [ws.cell(row=3, column=c).value for c in range(7, 13)]
    assert row3_vals[0] == "99.99.01.000-00000002"
    assert row3_vals[1] == "Ноутбук игровой кастомный"
    assert row3_vals[3] == "Исправлено"
    assert row3_vals[4] == "Пользователь"

    # Строка 7 (Excel row 9): Ошибка
    row9_vals = [ws.cell(row=9, column=c).value for c in range(7, 13)]
    assert row9_vals[0] is None or row9_vals[0] == ""
    assert row9_vals[3] == "Ошибка"

    # Строка 2 (Excel row 4): Проверка экранирования формулы (=1+1) в XLSX (ячейки имеют тип 's')
    assert ws.cell(row=4, column=2).data_type == "s"

    # 5.4 Проверка валидности через Calamine (быстрый парсер Rust)
    calamine_wb = CalamineWorkbook.from_filelike(io.BytesIO(res_xlsx.content))
    calamine_sheet = calamine_wb.get_sheet_by_name("Спецификация")
    calamine_data = calamine_sheet.to_python()
    assert len(calamine_data) == 101
    assert len(calamine_data[0]) == 12


def test_e2e_pipeline_csv_and_formula_injection(spec_100_csv: Path):
    """
    Тестирование конвейера для CSV с акцентом на защиту от Formula Injection
    и поддержку форматов выгрузки (CSV, JSON, SOURCE).
    """
    file_info = read_file(spec_100_csv)
    assert file_info.filename == "specification_100.csv"
    assert len(file_info.sheets) == 1

    sheet = file_info.sheets[0]
    headers = [c.header for c in sheet.columns]
    mapping = [ColumnMappingItem(index=c.index, role=c.suggested_role) for c in sheet.columns]

    parsed_rows = parse_rows(
        path=spec_100_csv,
        header_row=0,
        column_mapping=mapping,
    )
    assert len(parsed_rows) == 100

    # Создаем фиктивные элементы экспорта
    export_items = [
        ExportItem(
            row_index=r.row_index,
            raw_row=r.raw_row,
            ktru_code="99.99.01.000-00000001",
            ktru_name="Тестовый товар",
            confidence=0.95,
            status="confirmed" if r.row_index == 0 else "processed",
            source="prediction",
            flags=[],
        )
        for r in parsed_rows
    ]

    # --- Экспорт в CSV ---
    res_csv = export_job_results(
        items=export_items,
        headers=headers,
        export_format=ExportFormat.CSV,
        original_filename="specification_100.csv",
        allow_unreviewed=True,
    )

    assert res_csv.filename == "specification_100_ktru.csv"
    assert res_csv.media_type == "text/csv; charset=utf-8"

    # Проверка экранирования формул в CSV:
    # Строка 2 содержит '=1+1', '+2+2', '@SUM...'
    csv_text = res_csv.content.decode("utf-8-sig")
    lines = csv_text.splitlines()
    assert len(lines) == 101  # 1 заголовок + 100 строк

    row_with_formula = lines[3]  # заголовок + 3 строки
    assert "'=1+1" in row_with_formula
    assert "'+2+2" in row_with_formula
    assert "'@SUM" in row_with_formula

    # --- Экспорт в JSON ---
    res_json = export_job_results(
        items=export_items,
        headers=headers,
        export_format=ExportFormat.JSON,
        original_filename="specification_100.csv",
        allow_unreviewed=True,
    )
    assert res_json.filename == "specification_100_ktru.json"
    parsed_json = json.loads(res_json.content.decode("utf-8"))
    assert len(parsed_json) == 100
    assert parsed_json[0]["ktru_code"] == "99.99.01.000-00000001"
    assert parsed_json[0]["status"] == "confirmed"
    assert parsed_json[0]["status_label"] == "Подтверждено"

    # --- Экспорт в SOURCE (для CSV должен вернуть CSV) ---
    res_source = export_job_results(
        items=export_items,
        headers=headers,
        export_format=ExportFormat.SOURCE,
        original_filename="specification_100.csv",
        allow_unreviewed=True,
    )
    assert res_source.filename == "specification_100_ktru.csv"


def test_e2e_pipeline_1000_items(spec_1000_xlsx: Path):
    """
    Сквозной тест обработки максимальной спецификации на 1000 позиций (Критерий 3).
    Проверяет производительность этапов: read (< 2 с), parse, ML-заглушка и экспорт.
    """
    import time

    # 1. Быстрое чтение файла (NFR: < 2 секунд)
    t0 = time.perf_counter()
    file_info = read_file(spec_1000_xlsx)
    t_read = time.perf_counter() - t0
    assert t_read < 2.0, f"Чтение 1000 строк заняло слишком много времени: {t_read:.2f} с"
    assert file_info.sheets[0].row_count == 1000

    sheet = file_info.sheets[0]
    headers = [c.header for c in sheet.columns]
    mapping = [ColumnMappingItem(index=c.index, role=c.suggested_role) for c in sheet.columns]

    # 2. Разбор строк
    t0 = time.perf_counter()
    parsed_rows = parse_rows(
        path=spec_1000_xlsx,
        sheet_name=sheet.name,
        header_row=0,
        column_mapping=mapping,
    )
    t_parse = time.perf_counter() - t0
    assert len(parsed_rows) == 1000
    assert t_parse < 3.0, f"Разбор 1000 строк занял {t_parse:.2f} с"

    # 3. Пакетная классификация на ML-заглушке (с delay=0 для чистого теста логики)
    pipeline = StubPipeline()
    pipeline.delay = 0.0  # отключаем задержку сети/таймера
    raw_items = [
        RawItem(
            item_id=str(uuid4()),
            text=r.raw_text,
            columns={"name": r.raw_row[1], "quantity": r.raw_row[4], "unit": r.raw_row[5]},
        )
        for r in parsed_rows
    ]

    t0 = time.perf_counter()
    results = []
    chunk_size = 50
    for i in range(0, len(raw_items), chunk_size):
        results.extend(pipeline.classify_batch(raw_items[i : i + chunk_size], catalog_version_id="2026-10-25"))
    t_ml = time.perf_counter() - t0
    assert len(results) == 1000
    assert t_ml < 2.0

    # 4. Формирование позиций и экспорт в XLSX
    export_items = [
        ExportItem(
            row_index=r.row_index,
            raw_row=r.raw_row,
            ktru_code=res.predictions[0].ktru_code if res.predictions else None,
            ktru_name=res.predictions[0].name if res.predictions else None,
            confidence=res.predictions[0].confidence if res.predictions else None,
            status="failed" if res.error else "processed",
            source="prediction",
            flags=[],
        )
        for r, res in zip(parsed_rows, results, strict=False)
    ]

    t0 = time.perf_counter()
    res_xlsx = export_job_results(
        items=export_items,
        headers=headers,
        export_format=ExportFormat.XLSX,
        original_filename="specification_1000.xlsx",
        source_path=spec_1000_xlsx,
        sheet_name=sheet.name,
        allow_unreviewed=True,
        header_row=0,
    )
    t_export = time.perf_counter() - t0
    assert t_export < 5.0, f"Экспорт 1000 строк в XLSX занял {t_export:.2f} с"

    # 5. Проверка сформированного файла через Calamine
    calamine_wb = CalamineWorkbook.from_filelike(io.BytesIO(res_xlsx.content))
    calamine_sheet = calamine_wb.get_sheet_by_name("Спецификация")
    data = calamine_sheet.to_python()
    assert len(data) == 1001  # заголовок + 1000 строк
    assert len(data[0]) == 12  # 6 исходных + 6 добавленных

